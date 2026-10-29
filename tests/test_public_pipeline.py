import sqlite3

from docx import Document

from app.ingest.public import PublicCorpusPipeline, PublicDocument, PublicSource
from app.retrieval.embeddings import LocalHashEmbedder
from app.retrieval.store import LocalIndex


def test_public_pipeline_deduplicates_and_rolls_back(tmp_path) -> None:
    connection = sqlite3.connect(":memory:")
    index = LocalIndex(connection, 8)
    pipeline = PublicCorpusPipeline(connection, index, LocalHashEmbedder(8))
    path = tmp_path / "judgment.docx"
    document = Document()
    document.add_paragraph("Public judgment paragraph 1.")
    document.save(path)
    source = PublicSource("court", "Court judgments", "Open Government Licence")
    public_document = PublicDocument(
        "judgment-1",
        "court",
        path.name,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "https://example.test/judgment-1",
    )

    assert pipeline.ingest(source, public_document, path, "snapshot-1") == "indexed"
    assert pipeline.ingest(source, public_document, path, "snapshot-2") == "duplicate"
    assert len(index.search_lexical("public judgment", frozenset(), 3)) == 1

    pipeline.rollback("snapshot-1")
    assert index.search_lexical("public judgment", frozenset(), 3) == []
